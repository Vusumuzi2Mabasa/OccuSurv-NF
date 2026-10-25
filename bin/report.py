#!/usr/bin/env python3
"""Create a research report using exported evidence; no statistical refitting."""
import argparse,csv,json,platform
from pathlib import Path
import docx
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def read(path):
    with path.open() as f:return list(csv.DictReader(f))
def build(a):
    out=Path(a.out);out.mkdir(exist_ok=True);src=Path(a.analysis);figs=Path(a.figures)
    s=json.loads((src/'analysis.json').read_text());qc=json.loads((Path(a.clean)/'qc.json').read_text());methods=json.loads(Path(a.methods).read_text())
    d=Document();sec=d.sections[0];sec.page_width=Inches(8.5);sec.page_height=Inches(11);sec.top_margin=sec.bottom_margin=Inches(.7);sec.left_margin=sec.right_margin=Inches(.8)
    for name in ['Normal','Title','Subtitle','Heading 1','Heading 2']:
        st=d.styles[name];st.font.name='Calibri';st.font.color.rgb=RGBColor(0,0,0)
    d.styles['Normal'].font.size=Pt(11);d.styles['Normal'].paragraph_format.space_after=Pt(7)
    d.styles['Normal'].paragraph_format.line_spacing=1.08
    d.styles['Title'].font.size=Pt(25);d.styles['Heading 1'].font.size=Pt(17);d.styles['Heading 2'].font.size=Pt(13)
    def para(t):d.add_paragraph(t)
    def heading(t):d.add_heading(t,1)
    def page(t):d.add_page_break();heading(t)
    def table(headers,rows):
        t=d.add_table(rows=1,cols=len(headers));t.style='Light Shading Accent 1'
        for c,h in zip(t.rows[0].cells,headers):c.text=str(h)
        for row in rows:
            cells=t.add_row().cells
            for c,v in zip(cells,row):c.text=str(v)
        for row in t.rows:
            for c in row.cells:
                for p in c.paragraphs:
                    for r in p.runs:r.font.size=Pt(10)
        for cell in t.rows[0].cells:
            shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E8EDF0');cell._tc.get_or_add_tcPr().append(shade)
            for r in cell.paragraphs[0].runs:r.bold=True;r.font.color.rgb=RGBColor(0,0,0)
        repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
        d.add_paragraph()
    def image(name,caption,width=6.6):
        d.add_picture(str(figs/name),width=Inches(width));p=d.add_paragraph(caption);p.paragraph_format.space_after=Pt(9)
        for r in p.runs:r.font.size=Pt(10)
    def equation(i):
        e=next(e for e in methods if e['id']==i);d.add_picture(str(Path(a.equations)/(i+'.png')),width=Inches({'residual':4.4,'cox':5.8,'newton':5.8,'hr':5.5,'km':5.6}.get(i,6.6)));para(e['explanation'])
    d.add_heading('Occupational exposure and respiratory disease surveillance',0)
    para('A reproducible synthetic cohort analysis using Nextflow')
    para('Research software validation report | OccuSurv NF version 0.2.0 | Seed '+str(s['truth']['seed']))
    para('This study evaluates a reproducible workflow for exposure monitoring, disease surveillance and time-to-event analysis. All worker records and outcomes are simulated. The analysis demonstrates computational behaviour under stated assumptions and does not estimate disease burden or risk for NIOH or any real workforce.')
    heading('Principal findings')
    overall=read(src/'incidence.csv')[0];prev=read(src/'prevalence.csv')[0]
    para(f"After quality control, {s['workers']} workers remained. The incident cohort included {s['incident_cohort']} workers, {s['events']} diagnoses and {s['person_years']:.1f} person-years. The crude incidence rate was {float(overall['rate_per_1000']):.2f} per 1000 person-years. Regression status was {s['model']['status']}.")
    table(['Measure','Result'],[['Generated workers',s['truth']['n_workers']],['Eligible workers',s['workers']],['Baseline disease',f"{s['prevalent_cases']} / {s['workers']}"],['Incident cohort',s['incident_cohort']],['Incident events',s['events']],['Person years',f"{s['person_years']:.1f}"],['Complete case model events',s['model']['events']]])
    para('The pipeline separates prevalence, exposure summaries and incidence denominators. It preserves row-level QC reasons and skips regression when fewer than 30 complete-case events are available. The principal model estimates are evaluated against the assumed generating values, without requiring statistical significance or forcing interval coverage.')
    page('Simulation design and exposure history')
    para('Workers are assigned to twelve workplaces across construction, brick manufacture and stone fabrication. Age is uniform from 25 to 65 years. Employment duration is Uniform(3,30), capped at age minus 18. Smoking is Bernoulli with probability 0.35.')
    para('Historical concentration is lognormal, with industry medians 0.04, 0.07 and 0.12 mg/m3, workplace log SD 0.18 and worker log SD 0.35. Each worker has three baseline measurements with log SD 0.25 measurement error. Workplace effects influence exposure only; the model contains no unmeasured shared disease frailty.')
    equation('exposure');equation('hazard');equation('event')
    para('Baseline disease probability is expit(-3.4 + 0.5x + 0.3a). Incident times are generated among baseline disease-free workers conditional on their covariates. The exposure lag and hazard coefficients are illustrative assumptions, not empirically calibrated silicosis parameters. The model does not reconstruct lifetime disease onset.')
    page('Quality control and disease burden estimators')
    table(['QC measure','Count'],[['Input worker rows including duplicate',qc['workers_raw']],['Retained workers',qc['workers_retained']],['Valid follow-up records',qc['followup_retained']],['Measurements before QC',qc['measurements_raw']],['Measurements after QC',qc['measurements_retained']],['Audit entries including warnings',qc['audit_entries']]])
    para('Exact duplicate worker rows are collapsed; conflicting IDs fail validation. Invalid age, employment history and baseline status are excluded. Follow-up must have positive time no greater than eight years and a binary event. Measurements require valid IDs, nonnegative concentrations, acceptable sampling durations and mg/m3 units. Missing smoking is retained for descriptive analyses and excluded from regression. Audit entries may overlap and must not be added as unique exclusions.')
    equation('wilson');equation('rate');equation('km')
    page('Regression estimation and diagnostics')
    equation('cox');equation('newton');equation('hr');equation('residual')
    para('The Cox model uses complete cases and no automated variable selection. It assumes noninformative censoring and proportional hazards. Conditional independence is built into the simulation; the model-based standard errors should not be transported to a real clustered workforce without reassessing dependence. The event threshold is a guardrail, not a proof of adequate power.')
    page('Exposure and incidence results')
    image('exposure.png','Figure 1. Each point is the mean of valid baseline measurements for a worker. Boxes summarise the worker means; no occupational limit is imposed.')
    image('incidence.png','Figure 2. Industry-specific crude rates with conventional 95 percent Poisson intervals. These are not adjusted exposure effects.')
    para('Repeated monitoring measurements contribute to worker-level descriptive means. The regression instead uses the known simulated historical concentration. Measurement error in reconstructed occupational history remains a future extension.')
    page('Disease free survival')
    image('survival.png','Figure 3. Kaplan Meier survival with pointwise log-minus-log intervals. Prevalent cases are excluded. Curves are descriptive and unadjusted.')
    risk=read(src/'risk_table.csv');groups=sorted({r['industry'] for r in risk})
    table(['Industry','0 y','2 y','4 y','6 y','8 y'],[[g]+[next(r['at_risk'] for r in risk if r['industry']==g and float(r['time'])==t) for t in [0,2,4,6,8]] for g in groups])
    rates=read(src/'incidence.csv')
    table(['Industry','Events','Person years','Rate per 1000'],[[r['industry'],r['events'],f"{float(r['person_years']):.1f}",f"{float(r['rate_per_1000']):.2f}"] for r in rates])
    para('Risk sets count workers observed immediately before the stated time. Administrative censoring occurs at eight years. Later estimates have fewer workers at risk; uncertainty should be interpreted accordingly.')
    page('Adjusted associations')
    if s['model']['status']=='fitted':
        image('cox.png','Figure 4. Estimated hazard ratios and 95 percent Wald intervals compared with the assumed generating values. Crosses indicate simulation parameters, not external estimates.')
        rr=read(src/'cox.csv');table(['Term','HR','95 percent CI','P'],[[r['term'],f"{float(r['hr']):.3f}",f"{float(r['lower']):.3f} to {float(r['upper']):.3f}",f"{float(r['p']):.3g}"] for r in rr])
        image('ph_diagnostics.png','Figure 5. Schoenfeld residuals and LOWESS trends against log event time. These are exploratory diagnostics rather than a formal test.',6.6)
    else:para('The Cox model was not fitted: '+s['model']['reason']+'. No effect estimates or model diagnostics are manufactured.')
    ext=Path(a.longitudinal);ex=json.loads((ext/'summary.json').read_text());oldfigs=figs;figs=Path(a.longitudinal_figures)
    page('Time varying exposure and competing events')
    para('This extension generates new follow-up outcomes for the same eligible baseline disease-free workers. It is a separate scenario, not an additional follow-up phase and not a paired estimate of the effect of adding mortality. The preceding baseline reference results are retained unchanged.')
    para('Annual concentration equals historical concentration times independent lognormal noise with log SD 0.25, multiplied by 0.65 from year three. This planned reduction is illustrative, not an estimated causal intervention. Historical exposure enters the cumulative measure only once it is at least five years old.')
    equation('tv_exposure');equation('tv_hazard');equation('tv_risk')
    para('Within each annual interval, waiting time is exponential with the sum of the two hazard rates. If a first event occurs before the next boundary or independent censoring, its cause is selected with probability proportional to its rate. Follow-up stops at that event. Censoring is the minimum of eight years and Exponential(mean 25 years).')
    para('QC requires contiguous nonoverlapping intervals from baseline to terminal time, one terminal event at most, consistent cause coding, valid exposure schedules and no post-event observations. Only complete workers enter both cause-specific models. Each model is skipped below 30 events of its own cause.')
    page('Cumulative incidence with competing mortality')
    equation('aj')
    image('cumulative_incidence.png','Figure 6. Aalen-Johansen disease and death probabilities with pointwise bootstrap intervals. The dashed disease curve treats death as censoring and therefore does not estimate the observed disease probability.')
    endpoints=[r for r in read(ext/'cif_intervals.csv') if r['group']=='Overall' and float(r['time'])==8]
    table(['Cause at eight years','Probability','95 percent interval'],[[r['cause'],f"{float(r['cif']):.3f}",f"{float(r['lower']):.3f} to {float(r['upper']):.3f}"] for r in endpoints])
    para(f"The extension includes {ex['subjects']} workers, {ex['intervals']} at-risk intervals, {ex['disease']} diagnoses and {ex['deaths']} competing deaths over {ex['person_years']:.1f} person-years. Independent censoring and absence of residual workplace dependence are assumed. Bootstrap draws resample whole workers, never individual interval rows.")
    page('Exposure trajectories and industry probabilities')
    image('exposure_history.png','Figure 7. Medians among workers still under observation. Changes reflect both the planned schedules and changing risk-set composition; later latent exposure values are never used after an event.')
    image('cif_by_industry.png','Figure 8. Unadjusted diagnosis cumulative incidence by industry with pointwise bootstrap intervals. Industry comparisons are descriptive and incorporate competing mortality.')
    para('Full planned exposure histories are stored for auditing the simulation. They are explicitly latent schedules, not post-death measurements. The models use only rows with positive at-risk duration before the first event or censoring.')
    page('Time varying cause specific associations')
    image('time_varying_hr.png','Figure 9. Cause-specific Cox HRs with 95 percent model-based Wald intervals. Crosses show assumed generating HRs. An HR changes an instantaneous rate among event-free workers, not an absolute cumulative risk.')
    rows=[]
    labels={'log2_1plus_lagged':'Lagged exposure','age_per_decade':'Age per decade','smoker':'Smoker'}
    for cause in [1,2]:
        f=ext/f'cox_cause{cause}.csv'
        if f.exists():
            for r in read(f):rows.append(['Diagnosis' if cause==1 else 'Death',labels[r['term']],f"{float(r['hr']):.3f}",f"{float(r['lower']):.3f} to {float(r['upper']):.3f}"])
    if rows:table(['Cause','Covariate','HR','95 percent CI'],rows)
    else:para('Both cause-specific models were skipped because too few events were observed.')
    para('The split-record Cox analysis uses interval entry times and a single active record per worker at each event. Model-based information provides uncertainty under independent worker event processes; repeated rows are not treated as recurrent events. Both competing events remove a worker from subsequent risk sets. Missing smoking remains missing for estimation, although a latent smoking value is simulated to generate outcomes.')
    para('No Fine-Gray model, causal effect of exposure reduction or individual risk prediction is claimed. The bootstrap uses 200 replicates to keep the demonstration small; more replicates and cluster-aware analysis would be needed for a real workforce.')
    figs=oldfigs
    page('Interpretation limitations and reproducibility')
    para('This simulation verifies data flow and numerical execution for a small dataset. Agreement with generating values in one run does not establish estimator unbiasedness, confidence interval coverage or suitability for surveillance decisions. Statistical significance is not a validation criterion.')
    para('The baseline reference scenario omits competing mortality and time-varying exposure; the extension adds these explicitly. Both scenarios omit healthy worker selection, diagnostic misclassification and informative loss to follow-up. The five-year lag is a declared modelling choice. Twelve workplaces are not sufficient to justify complex hierarchical inference without further investigation. The independent nature of the project is explicit; it uses no institutional patient data and is not endorsed by NIOH.')
    para('Nextflow runs simulation, QC, statistical analysis, figures and this report as dependent processes. Outputs include CSV tables, QC audit, hashes, assumptions, software versions and execution trace. Resume behaviour and a low-event smoke run are checked separately in the repository validation record.')
    table(['Component','Version'],[[k,v] for k,v in s['versions'].items()]+[['Report Python',platform.python_version()],['python docx',docx.__version__]])
    heading('Method references')
    for t in ['Nextflow installation and execution: https://docs.seqera.io/nextflow/', 'Statsmodels proportional hazards regression: https://www.statsmodels.org/stable/duration.html', 'Statsmodels PHReg model: https://www.statsmodels.org/stable/generated/statsmodels.duration.hazard_regression.PHReg.html', 'Statistical equations, assumptions and exact implementation are provided in docs/METHODS.md and the versioned source scripts.']:para(t)
    d.core_properties.title='Occupational exposure and respiratory disease surveillance';d.core_properties.author='OccuSurv NF';d.core_properties.subject='Synthetic cohort methodology and computational results'
    for element in [d._element]+[style.element for style in d.styles]:
        for border in list(element.iter(qn('w:pBdr'))):border.getparent().remove(border)
    d.save(out/'OccuSurv_report.docx')
    (out/'report_versions.json').write_text(json.dumps(dict(python=platform.python_version(),python_docx=docx.__version__),indent=2))
if __name__=='__main__':
    p=argparse.ArgumentParser()
    for name in ['analysis','clean','figures','methods','equations','longitudinal','longitudinal_figures']:p.add_argument('--'+name,required=True)
    p.add_argument('--out',default='report');build(p.parse_args())
